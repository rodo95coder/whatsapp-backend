from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

OUT = Path('docs_actualizados'); OUT.mkdir(exist_ok=True)

API = [
('Visión general', 'API REST multi-sesión para WhatsApp Web basada en Node.js, Express y WPPConnect. Cada companyId usa un perfil Chromium aislado.'),
('Autenticación', 'Todas las rutas /api requieren Authorization: Bearer <GLOBAL_TOKEN> o el header X-Global-Token. No documente ni comparta el valor real del token.'),
('Sesiones', 'POST /api/whatsapp/init-session recibe {"companyId":"empresa_demo"}. GET /api/whatsapp/{companyId}/status devuelve status, reason, updatedAt, generationId y recoverable. GET /api/whatsapp/{companyId}/qr devuelve un PNG QR de alto contraste mientras status sea QR_REQUIRED.'),
('Estados', 'IDLE, CONNECTING, QR_REQUIRED, CONNECTED, RECONNECTING, DISCONNECTED, SHUTTING_DOWN, FAILED, INITIALIZATION_TIMEOUT y NOT_FOUND. Solo CONNECTED permite envíos.'),
('Mensajes', 'POST /api/whatsapp/{companyId}/send recibe numbers (array), text opcional, y opcionalmente base64File o fileUrl con fileName. La respuesta incluye resultado por número. sent_unconfirmed significa que WPPConnect no confirmó el resultado; no debe reintentarse automáticamente.'),
('Confirmaciones', 'Cuando WPPConnect devuelve messageId, GET /api/whatsapp/{companyId}/messages/{messageId}/status consulta el ACK temporal. Estados: pending, server, delivered, read, played o error. El historial está en memoria y expira; un reinicio lo elimina.'),
('Recuperación', 'POST /api/whatsapp/{companyId}/force-reset recibe deleteAuth, deleteSessionMetadata y restart. Use primero false, false, true. deleteAuth=true borra el perfil y exige QR nuevo. POST /logout elimina autenticación y metadatos.'),
('Webhooks', 'POST /api/webhook/{companyId}/webhook configura {"url":"https://..."}. Los ACK rastreados emiten event=message.ack. Las URLs privadas, localhost y destinos SSRF se bloquean.'),
('Seguridad y límites', 'companyId admite letras, números, guion y underscore. Las descargas externas se validan contra SSRF. El tracker de ACK usa TTL, límites por sesión/globales y purga por presión de memoria.'),
]

OPS = [
('Requisitos', 'Node.js 20+, Chromium/Google Chrome compatible, PM2, Nginx y un usuario de servicio sin privilegios root. Instale dependencias con npm install.'),
('Configuración', 'Cree .env con PORT, GLOBAL_TOKEN, LOG_LEVEL, SESSIONS_PATH=./userSessions, TEMP_PATH=./temp y PUPPETEER_EXECUTABLE_PATH. Nunca suba .env al repositorio.'),
('Variables operativas', 'QR_TIMEOUT_MS=180000, INIT_SESSION_TIMEOUT_MS=180000, BROWSER_CLOSE_TIMEOUT_MS=10000, PUPPETEER_PROTOCOL_TIMEOUT_MS=60000, SESSION_WATCHDOG_INTERVAL_MS=30000.'),
('Capacidad', 'No existe un límite duro de sesiones. El monitor registra session.capacity.warning cuando la memoria libre del sistema o RSS de Node supera umbrales. Ajuste SESSION_CAPACITY_WARN_FREE_MEMORY_MB, SESSION_CAPACITY_WARN_PROCESS_RSS_MB y SESSION_CAPACITY_WARN_COOLDOWN_MS.'),
('PM2 y Nginx', 'Inicie con pm2 start ecosystem.config.js y use pm2 logs whatsapp-backend. El archivo nginx/whatsapp-backend.conf es la referencia de proxy. Proteja el servidor con HTTPS y no exponga el puerto interno si Nginx atiende el tráfico.'),
('Operación', 'Para QR: init-session, consultar status hasta QR_REQUIRED, consultar /qr y escanear. No vuelva a llamar init-session mientras espera QR. Para una sesión bloqueada use force-reset conservando autenticación antes de borrar perfiles.'),
('Observabilidad', 'Revise logs por companyId, generation y reason. INITIALIZATION_TIMEOUT señala arranque bloqueado; QR_TIMEOUT señala que no se escaneó a tiempo. Los ACK y advertencias de capacidad quedan en logs.'),
('Backup y actualización', 'Respalde userSessions con el servicio detenido o con perfiles consistentes. Antes de desplegar ejecute node --test y eslint. Actualice con git pull, npm install y pm2 restart whatsapp-backend.'),
]

def add_docx(title, sections, filename):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75)
    styles=d.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10)
    p=d.add_paragraph(); r=p.add_run(title); r.bold=True; r.font.name='Arial'; r.font.size=Pt(22)
    d.add_paragraph('Actualizado para la versión multi-sesión actual - agosto de 2026.')
    for heading, body in sections:
        d.add_heading(heading, level=1); d.add_paragraph(body)
    d.add_heading('Ejemplo de estado', level=1); d.add_paragraph('{"success":true,"status":"CONNECTED","reason":null,"recoverable":true}')
    d.save(OUT/filename)

def add_pdf(title, sections, filename):
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle('H', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=14, spaceBefore=12, spaceAfter=6))
    story=[Paragraph(title, styles['Title']), Paragraph('Actualizado para la versión multi-sesión actual - agosto de 2026.', styles['Normal']), Spacer(1,12)]
    for heading, body in sections:
        story += [Paragraph(heading, styles['H']), Paragraph(body, styles['BodyText'])]
    story += [Paragraph('Ejemplo de estado', styles['H']), Paragraph('{"success":true,"status":"CONNECTED","reason":null,"recoverable":true}', styles['Code'])]
    SimpleDocTemplate(str(OUT/filename), pagesize=letter, leftMargin=.75*inch, rightMargin=.75*inch, topMargin=.7*inch, bottomMargin=.7*inch).build(story)

add_docx('Referencia API - WhatsApp Backend', API, 'documentacion_WhatsappBackend_actualizada.docx')
add_docx('Manual de Desarrollo, Operación y Despliegue - WhatsApp Backend', OPS, 'manual_backend_completo_actualizado.docx')
for name in ['whatsapp_api_doc_actualizado.pdf','documentacion_WhatsappBackend_actualizada.pdf']:
    add_pdf('Referencia API - WhatsApp Backend', API, name)
for name in ['manual_backend_despliegue_actualizado.pdf','manual_backend_extendido_actualizado.pdf','manual_backend_actualizado.pdf','manual_backend_completo_actualizado.pdf']:
    add_pdf('Manual de Operación y Despliegue - WhatsApp Backend', OPS, name)
